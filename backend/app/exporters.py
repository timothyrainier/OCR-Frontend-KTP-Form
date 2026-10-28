"""Ekspor hasil OCR: CSV, Excel (.xlsx), JSON, SQL dump, dan kirim langsung ke database."""
import csv
import io
import json
from datetime import datetime

from .ktp_parser import KTP_FIELDS

META_COLS = [("id", "ID"), ("filename", "Nama File"), ("doc_type", "Jenis"), ("status", "Status"),
             ("confidence", "Akurasi"), ("attempts", "Jumlah Scan"), ("edited", "Dikoreksi"),
             ("verified", "Terverifikasi"), ("created_at", "Waktu Scan")]


def _flat(doc: dict) -> dict:
    """{label: value} dari item hasil OCR."""
    return {i["label"]: i.get("value", "") for i in (doc.get("fields") or {}).get("items", [])}


def _meta(doc):
    out = {}
    for k, label in META_COLS:
        v = doc.get(k)
        if k == "confidence" and v is not None:
            v = round(v * 100, 1)
        if k in ("edited", "verified"):
            v = "Ya" if v else "Tidak"
        if k == "doc_type":
            v = "KTP" if v == "ktp" else "Form"
        out[label] = v
    return out


def _columns(docs):
    ktp_cols = [lbl for _, lbl, _ in KTP_FIELDS]
    extra = []
    for d in docs:
        for lbl in _flat(d):
            if lbl not in ktp_cols and lbl not in extra:
                extra.append(lbl)
    used_ktp = [c for c in ktp_cols if any(d["doc_type"] == "ktp" for d in docs)]
    return [lbl for _, lbl in META_COLS] + used_ktp + extra


def to_csv(docs) -> bytes:
    cols = _columns(docs)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for d in docs:
        w.writerow({**_meta(d), **_flat(d)})
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM agar Excel membaca UTF-8


def to_json(docs) -> bytes:
    out = []
    for d in docs:
        out.append({
            "id": d["id"], "filename": d["filename"], "doc_type": d["doc_type"], "status": d["status"],
            "confidence": d.get("confidence"), "attempts": d.get("attempts"), "edited": d.get("edited"),
            "verified": d.get("verified"), "created_at": d["created_at"], "updated_at": d["updated_at"],
            "data": {i["key"]: i.get("value", "") for i in (d.get("fields") or {}).get("items", [])},
            "fields": (d.get("fields") or {}).get("items", []),
            "tables": (d.get("fields") or {}).get("tables", []),
            "warnings": d.get("warnings", []), "error": d.get("error_message"),
        })
    return json.dumps({"exported_at": datetime.now().isoformat(timespec="seconds"),
                       "count": len(out), "documents": out}, ensure_ascii=False, indent=2).encode("utf-8")


def to_xlsx(docs) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F5FAD")
    head_font = Font(bold=True, color="FFFFFF")

    def write_sheet(ws, cols, rows):
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for i, col in enumerate(cols, start=1):
            width = max([len(str(col))] + [len(str(r.get(col, "") or "")) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 50)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    ktp = [d for d in docs if d["doc_type"] == "ktp"]
    form = [d for d in docs if d["doc_type"] == "form"]
    meta_cols = [lbl for _, lbl in META_COLS]

    ws = wb.active
    ws.title = "KTP"
    write_sheet(ws, meta_cols + [lbl for _, lbl, _ in KTP_FIELDS], [{**_meta(d), **_flat(d)} for d in ktp])

    ws2 = wb.create_sheet("Form")
    form_labels = []
    for d in form:
        for lbl in _flat(d):
            if lbl not in form_labels:
                form_labels.append(lbl)
    write_sheet(ws2, meta_cols + form_labels, [{**_meta(d), **_flat(d)} for d in form])

    table_rows = []
    for d in form:
        for t in (d.get("fields") or {}).get("tables", []):
            for row in t.get("rows", []):
                rec = {"ID Dokumen": d["id"], "Nama File": d["filename"], "Tabel": t.get("title", "")}
                for i, cell in enumerate(row):
                    hdr = t["headers"][i] if i < len(t.get("headers", [])) else f"Kolom {i + 1}"
                    rec[hdr] = cell
                table_rows.append(rec)
    if table_rows:
        cols = []
        for r in table_rows:
            cols += [k for k in r if k not in cols]
        write_sheet(wb.create_sheet("Tabel Form"), cols, table_rows)

    ws3 = wb.create_sheet("Ringkasan")
    ok = sum(d["status"] == "success" for d in docs)
    for row in [("Diekspor pada", datetime.now().strftime("%d-%m-%Y %H:%M")), ("Jumlah dokumen", len(docs)),
                ("KTP", len(ktp)), ("Form", len(form)), ("Berhasil", ok), ("Gagal", len(docs) - ok)]:
        ws3.append(row)
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------- database
def _ktp_record(d):
    vals = {i["key"]: i.get("value", "") for i in d["fields"].get("items", [])}
    return {"source_doc_id": d["id"], "filename": d["filename"], "confidence": d.get("confidence"),
            "verified": int(bool(d.get("verified"))), "scanned_at": d["created_at"],
            **{k: vals.get(k, "") for k, _, _ in KTP_FIELDS}}


def _sql_literal(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def to_sql(docs) -> bytes:
    """File .sql (INSERT) yang bisa diimpor ke PostgreSQL/MySQL/SQLite."""
    ktp_cols = ["source_doc_id", "filename", "confidence", "verified", "scanned_at"] + [k for k, _, _ in KTP_FIELDS]
    lines = ["-- Ekspor hasil OCR " + datetime.now().isoformat(timespec="seconds"),
             "CREATE TABLE IF NOT EXISTS ocr_ktp (" + ", ".join(
                 ["source_doc_id INTEGER PRIMARY KEY", "filename VARCHAR(255)", "confidence REAL",
                  "verified INTEGER", "scanned_at VARCHAR(25)"] + [f"{k} VARCHAR(255)" for k, _, _ in KTP_FIELDS]) + ");",
             "CREATE TABLE IF NOT EXISTS ocr_form_fields (source_doc_id INTEGER, filename VARCHAR(255), "
             "field_label VARCHAR(255), field_value TEXT, confidence REAL, scanned_at VARCHAR(25));", ""]
    for d in docs:
        if d["status"] != "success":
            continue
        if d["doc_type"] == "ktp":
            r = _ktp_record(d)
            lines.append(f"INSERT INTO ocr_ktp ({', '.join(ktp_cols)}) VALUES ({', '.join(_sql_literal(r[c]) for c in ktp_cols)});")
        else:
            for i in d["fields"].get("items", []):
                lines.append("INSERT INTO ocr_form_fields VALUES (" + ", ".join(_sql_literal(x) for x in (
                    d["id"], d["filename"], i["label"], i.get("value", ""), i.get("confidence"), d["created_at"])) + ");")
    return "\n".join(lines).encode("utf-8")


def push_to_database(docs, url: str) -> dict:
    """Simpan/upsert hasil ke database tujuan (EXPORT_DATABASE_URL) memakai SQLAlchemy."""
    from sqlalchemy import Column, Float, Integer, MetaData, String, Table, Text, create_engine, delete, insert

    engine = create_engine(url)
    md = MetaData()
    ktp = Table("ocr_ktp", md, Column("source_doc_id", Integer, primary_key=True),
                Column("filename", String(255)), Column("confidence", Float), Column("verified", Integer),
                Column("scanned_at", String(25)), *[Column(k, String(255)) for k, _, _ in KTP_FIELDS])
    form = Table("ocr_form_fields", md, Column("id", Integer, primary_key=True, autoincrement=True),
                 Column("source_doc_id", Integer, index=True), Column("filename", String(255)),
                 Column("field_label", String(255)), Column("field_value", Text), Column("confidence", Float),
                 Column("scanned_at", String(25)))
    md.create_all(engine)
    n_ktp = n_form = skipped = 0
    with engine.begin() as conn:
        for d in docs:
            if d["status"] != "success":
                skipped += 1
                continue
            if d["doc_type"] == "ktp":
                conn.execute(delete(ktp).where(ktp.c.source_doc_id == d["id"]))
                conn.execute(insert(ktp).values(**_ktp_record(d)))
                n_ktp += 1
            else:
                conn.execute(delete(form).where(form.c.source_doc_id == d["id"]))
                for i in d["fields"].get("items", []):
                    conn.execute(insert(form).values(source_doc_id=d["id"], filename=d["filename"],
                                                     field_label=i["label"], field_value=i.get("value", ""),
                                                     confidence=i.get("confidence"), scanned_at=d["created_at"]))
                n_form += 1
    safe_url = engine.url.render_as_string(hide_password=True)
    return {"ktp": n_ktp, "form": n_form, "skipped_failed": skipped, "target": safe_url}
